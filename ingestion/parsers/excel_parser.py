from langchain_core.documents import Document
from openpyxl import load_workbook


def parse_excel(file_path: str) -> list[Document]:
    workbook = load_workbook(
        filename=file_path,
        data_only=True,
        read_only=True,
    )

    documents = []

    for worksheet in workbook.worksheets:
        rows = worksheet.iter_rows(values_only=True)

        lines = []

        for row in rows:
            values = [str(value) if value is not None else "" for value in row]

            if any(value.strip() for value in values):
                lines.append(" | ".join(values))

        sheet_text = "\n".join(lines)

        if sheet_text.strip():
            documents.append(
                Document(
                    page_content=sheet_text,
                    metadata={
                        "source": file_path,
                        "file_type": "xlsx",
                        "sheet_name": worksheet.title,
                    },
                )
            )

    return documents
