from docx import Document as DocxDocument
from langchain_core.documents import Document


def parse_docx(file_path: str) -> list[Document]:
    doc = DocxDocument(file_path)

    text_parts = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text_parts.append(paragraph.text)

    text = "\n".join(text_parts)

    return [
        Document(
            page_content=text,
            metadata={
                "source": file_path,
                "file_type": "docx",
            },
        )
    ]
