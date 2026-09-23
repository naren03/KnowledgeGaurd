from langchain_core.documents import Document
from pptx import Presentation


def parse_pptx(file_path: str) -> list[Document]:
    presentation = Presentation(file_path)
    documents = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        text_parts = []

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text_parts.append(shape.text)

        slide_text = "\n".join(text_parts)

        if slide_text.strip():
            documents.append(
                Document(
                    page_content=slide_text,
                    metadata={
                        "source": file_path,
                        "file_type": "pptx",
                        "slide_number": slide_number,
                    },
                )
            )

    return documents
